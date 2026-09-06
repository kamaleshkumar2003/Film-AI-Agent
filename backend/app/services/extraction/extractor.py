import io
from pathlib import Path
from typing import Tuple
from pypdf import PdfReader
import docx
from app.core.exceptions import ScreenplayExtractionError
from app.services.extraction.page_tracker import PageContent, ExtractedDocument

class DocumentExtractor:
    @staticmethod
    def extract(file_bytes: bytes, filename: str) -> ExtractedDocument:
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return DocumentExtractor._extract_pdf(file_bytes)
        elif ext in [".docx", ".doc"]:
            return DocumentExtractor._extract_docx(file_bytes)
        elif ext in [".txt", ".fountain"]:
            return DocumentExtractor._extract_txt(file_bytes)
        else:
            raise ScreenplayExtractionError(f"Unsupported screenplay format: '{ext}'. Supported: .pdf, .docx, .txt, .fountain")

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> ExtractedDocument:
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            pages: list[PageContent] = []
            full_text_parts: list[str] = []
            current_char = 0

            for idx, page in enumerate(reader.pages):
                page_num = idx + 1
                page_text = page.extract_text() or ""
                # Normalize line endings
                page_text = page_text.replace("\r\n", "\n").replace("\r", "\n")
                
                char_start = current_char
                char_end = char_start + len(page_text)
                current_char = char_end + 1  # accounting for joined newline
                
                pages.append(PageContent(
                    page_number=page_num,
                    text=page_text,
                    char_start=char_start,
                    char_end=char_end
                ))
                full_text_parts.append(page_text)

            full_text = "\n".join(full_text_parts)
            if not full_text.strip():
                raise ScreenplayExtractionError("The uploaded PDF does not contain extractable text (it may be a scanned image).")

            return ExtractedDocument(
                full_text=full_text,
                pages=pages,
                page_count=len(pages)
            )
        except ScreenplayExtractionError:
            raise
        except Exception as e:
            raise ScreenplayExtractionError(f"Failed to extract text from PDF: {str(e)}")

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> ExtractedDocument:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs]
            full_text = "\n".join(paragraphs).replace("\r\n", "\n").replace("\r", "\n")
            
            if not full_text.strip():
                raise ScreenplayExtractionError("The uploaded DOCX file is empty.")

            # Approximate pages at ~55 lines per screenplay page
            lines = full_text.split("\n")
            lines_per_page = 55
            pages: list[PageContent] = []
            current_char = 0
            
            for page_idx in range(0, max(1, len(lines)), lines_per_page):
                page_num = (page_idx // lines_per_page) + 1
                chunk_lines = lines[page_idx : page_idx + lines_per_page]
                chunk_text = "\n".join(chunk_lines)
                
                char_start = current_char
                char_end = char_start + len(chunk_text)
                current_char = char_end + 1
                
                pages.append(PageContent(
                    page_number=page_num,
                    text=chunk_text,
                    char_start=char_start,
                    char_end=char_end
                ))

            return ExtractedDocument(
                full_text=full_text,
                pages=pages,
                page_count=len(pages)
            )
        except ScreenplayExtractionError:
            raise
        except Exception as e:
            raise ScreenplayExtractionError(f"Failed to extract text from DOCX: {str(e)}")

    @staticmethod
    def _extract_txt(file_bytes: bytes) -> ExtractedDocument:
        try:
            # Try utf-8 first, fallback to latin-1
            try:
                full_text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                full_text = file_bytes.decode("latin-1")

            full_text = full_text.replace("\r\n", "\n").replace("\r", "\n")
            if not full_text.strip():
                raise ScreenplayExtractionError("The uploaded text file is empty.")

            lines = full_text.split("\n")
            lines_per_page = 55
            pages: list[PageContent] = []
            current_char = 0

            for page_idx in range(0, max(1, len(lines)), lines_per_page):
                page_num = (page_idx // lines_per_page) + 1
                chunk_lines = lines[page_idx : page_idx + lines_per_page]
                chunk_text = "\n".join(chunk_lines)
                
                char_start = current_char
                char_end = char_start + len(chunk_text)
                current_char = char_end + 1
                
                pages.append(PageContent(
                    page_number=page_num,
                    text=chunk_text,
                    char_start=char_start,
                    char_end=char_end
                ))

            return ExtractedDocument(
                full_text=full_text,
                pages=pages,
                page_count=len(pages)
            )
        except ScreenplayExtractionError:
            raise
        except Exception as e:
            raise ScreenplayExtractionError(f"Failed to extract text from TXT: {str(e)}")
